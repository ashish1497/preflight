"""Build the demo/test workbook in the judges' template layout, with planted rulebook errors.

Run: python scripts/make_sample.py  (needs openpyxl)
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

GREY = PatternFill("solid", fgColor="D9D9D9")

B_HEAD = ["Booking ID", "Seller name", "Origin city", "Origin country", "Destination city", "Destination country",
          "Destination type", "Ready to ship date", "Seller priority", "Pallets", "Cartons", "Total declared value",
          "Currency", "Total gross weight (kg)", "Delivery appointment booked"]
I_HEAD = ["Booking ID", "SKU", "Description", "HS code", "Quantity (units)", "Unit price", "Units per carton",
          "Cartons", "Unit weight (kg)", "Country of origin"]

EXAMPLE_B = ["BK-001", "Northfield Goods", "Manchester", "United Kingdom", "Frankfurt", "Germany", "Warehouse",
             "05-11-2026", "Cheapest", 3, 90, 5400, "GBP", 620, "Yes"]
EXAMPLE_I = ["BK-001", "MUG-01", "Ceramic coffee mug, 350 ml", 691200, 1080, 5.00, 12, 90, 0.5, "United Kingdom"]

UK = "United Kingdom"
BOOKINGS = [
    # Clean, low value, Balanced -> economy wins
    ["BK-101", "Pennine Apparel", "Leeds", UK, "Paris", "France", "Business address", "06-11-2026", "Balanced", 2, 60, 6240, "GBP", 450, "No"],
    # R-03 (notebooks 25 x 38 != 1000) and R-04 (60 vs 62 cartons)
    ["BK-102", "Inkwell Stationery", "London", UK, "Berlin", "Germany", "Warehouse", "09-11-2026", "Fastest", 3, 60, 4420, "GBP", 420, "Yes"],
    # R-05: 1700 kg on 2 pallets = 850 kg/pallet, over every carrier's limit
    ["BK-103", "Forge & Hearth", "Birmingham", UK, "Rotterdam", "Netherlands", "Warehouse", "10-11-2026", "Balanced", 2, 50, 5400, "GBP", 1700, "Yes"],
    # R-01 (lost leading zero + missing HS) and R-09 (blank origin)
    ["BK-104", "Northern Roast Co", "Manchester", UK, "Hamburg", "Germany", "Warehouse", "11-11-2026", "Most reliable", 2, 60, 5760, "GBP", 640, "Yes"],
    # R-06 (power bank, perfume) and R-02 (9500 vs 11000)
    ["BK-105", "Gadget Garage", "London", UK, "Amsterdam", "Netherlands", "Business address", "12-11-2026", "Cheapest", 1, 30, 9500, "GBP", 300, "No"],
    # R-07 (200 kg < 240 kg goods), R-08 (warehouse, no slot), R-10 (date in the past)
    ["BK-106", "Dales Wool Mill", "Leeds", UK, "Lyon", "France", "Warehouse", "10-10-2026", "Fastest", 1, 20, 4400, "GBP", 200, "No"],
    # Clean, but 480 kg/pallet rules out express (400 kg limit) for a Fastest seller
    ["BK-107", "Midland Tile Supply", "Birmingham", UK, "Frankfurt", "Germany", "Warehouse", "13-11-2026", "Fastest", 2, 80, 5200, "GBP", 960, "Yes"],
    # Clean, high value, Balanced -> standard wins (time in transit costs more)
    ["BK-108", "Thames Horology", "London", UK, "Berlin", "Germany", "Warehouse", "16-11-2026", "Balanced", 3, 20, 60000, "GBP", 300, "Yes"],
    # D-01: description "Gifts" too vague for customs
    ["BK-109", "Little Treasures Ltd", "Manchester", UK, "Paris", "France", "Business address", "17-11-2026", "Cheapest", 1, 20, 3000, "GBP", 200, "No"],
]

ITEMS = [
    ["BK-101", "TEE-WHT-M", "Men's cotton T-shirt, white, size M", 61091000, 800, 6.00, 20, 40, 0.2, "India"],
    # Clean, but made in China: EU anti-dumping duty applies (shown by the customs-by-product check)
    ["BK-101", "MUG-STN-30", "Stoneware coffee mug, 300 ml", 69120023, 480, 3.00, 24, 20, 0.4, "China"],
    ["BK-102", "NB-A5-RUL", "A5 ruled notebook, 80 pages, paper cover", 48202000, 1000, 2.50, 25, 38, 0.3, UK],
    ["BK-102", "PEN-BLK-10", "Black ballpoint pen", 96081010, 2400, 0.80, 100, 24, 0.02, "China"],
    ["BK-103", "SKL-CI-26", "Cast iron skillet, 26 cm, uncoated", 73239100, 300, 18.00, 6, 50, 5.0, "China"],
    ["BK-104", "COF-1KG-VN", "Roasted coffee beans, 1 kg bag", 90121, 400, 12.00, 10, 40, 1.05, "Vietnam"],
    ["BK-104", "MUG-ESP", "Stoneware espresso cup, 90 ml", None, 240, 4.00, 12, 20, 0.4, None],
    ["BK-105", "PB-10K", "Power bank 10000 mAh, USB-C", 85076000, 500, 9.00, 50, 10, 0.25, "China"],
    ["BK-105", "EDP-50", "Eau de parfum, 50 ml spray bottle", 33030010, 300, 15.00, 30, 10, 0.2, "Turkey"],
    # 6-digit code: passes the length check, but isn't a full listed code (pick-list offered)
    ["BK-105", "CASE-SLM", "Silicone phone case", 392690, 1000, 2.00, 100, 10, 0.05, "China"],
    ["BK-106", "THR-WOOL", "Wool throw blanket, 130 x 170 cm", 63012010, 200, 22.00, 10, 20, 1.2, "Turkey"],
    ["BK-107", "TILE-30", "Glazed ceramic floor tile, 30 x 30 cm", 69072100, 800, 6.50, 10, 80, 1.1, "Turkey"],
    ["BK-108", "WATCH-AUTO", "Mechanical wristwatch, stainless steel case", 91022100, 400, 150.00, 20, 20, 0.3, "Germany"],
    ["BK-109", "GIFT-BOX", "Gifts", 95030021, 500, 6.00, 25, 20, 0.3, "China"],
]


def sheet(ws, title_text, head, example, rows):
    ws.append([title_text])
    ws.append(head)
    for c in ws[2]:
        c.font = Font(bold=True)
    ws.append(example)
    for c in ws[3]:
        c.fill = GREY
    for r in rows:
        ws.append(r)


wb = Workbook()
b = wb.active
b.title = "Bookings"
sheet(b, "Bookings — one row per record. Row 3 is an example (grey); enter data from row 4.", B_HEAD, EXAMPLE_B, BOOKINGS)
i = wb.create_sheet("Items")
sheet(i, "Items — one row per record. Row 3 is an example (grey); enter data from row 4.", I_HEAD, EXAMPLE_I, ITEMS)
wb.save("public/sample/ShipCheck_sample_bookings.xlsx")
print("ok", len(BOOKINGS), "bookings", len(ITEMS), "items")
